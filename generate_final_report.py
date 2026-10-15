"""Final project report (BCSE307P Innovative Project) -> Compiler_Design_Project_Final_Report.docx

    python generate_final_report.py              # figures + docx
    python generate_final_report.py --no-figures # docx only (reuse docs/report_figures/*.png)

Fills the VIT report template (title page + the template's seven headings) with the
900-program results.  Charts are drawn from data/benchmark_dataset.csv; terminal
figures are rendered from real command output captured into docs/report_figures/term_*.txt.
Needs: python-docx, matplotlib.
"""
import argparse
import math
import os
import re
from collections import defaultdict

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch

from src.minic.harness.report import summarize

ROOT = os.path.dirname(os.path.abspath(__file__))
FIG = os.path.join(ROOT, "docs", "report_figures")
DATASET = os.path.join(ROOT, "data", "benchmark_dataset.csv")
TEMPLATE = os.path.join(os.path.expanduser("~"), "Downloads", "Project - Report Template.docx")
OUT = os.path.join(ROOT, "Compiler_Design_Project_Final_Report.docx")

PASS_BITS = [(1, "CF"), (2, "DCE"), (4, "CSE"), (8, "LICM"), (16, "SR"), (32, "LU")]

# Strategy bake-off on the 900-program dataset (run_experiment.py output, 5-fold
# GroupKFold).  margin 0.0 = plain argmax.  (reco speedup, regression %, capture %)
STRATEGY = {
    "Random Forest": {0.0: (1.21, 12, 64), 0.03: (1.21, 10, 64), 0.05: (1.21, 9, 63),
                      0.08: (1.21, 8, 62), 0.12: (1.20, 4, 60), 0.18: (1.18, 3, 55)},
    "HistGradient Boosting": {0.0: (1.21, 14, 63), 0.03: (1.21, 10, 63), 0.05: (1.21, 9, 62),
                              0.08: (1.20, 6, 60), 0.12: (1.19, 4, 58), 0.18: (1.18, 2, 53)},
    "Extra Trees": {0.0: (1.21, 13, 62), 0.03: (1.21, 12, 63), 0.05: (1.21, 11, 62),
                    0.08: (1.20, 9, 61), 0.12: (1.20, 6, 59), 0.18: (1.19, 3, 56)},
}

# MODEL_COMPARISON.md: (input, reco @ best margin under 10% cap, capture %, regress %,
# MAE, size, train s/fold, uncapped reco, uncapped capture %, uncapped regress %)
MODELS = [
    ("HistGradient Boosting", "features", 1.209, 63, 10, 0.105, "2.0 MB", 1.6, 1.217, 65, 14),
    ("Random Forest", "features", 1.206, 62, 10, 0.107, "612.9 MB", 5.7, 1.209, 63, 13),
    ("1D-CNN", "TAC sequence", 1.197, 59, 8, 0.131, "365 KB", 43.4, 1.209, 63, 13),
    ("Linear (Ridge)", "features", 1.186, 56, 5, 0.136, "2 KB", 0.0, 1.205, 62, 14),
    ("1D-CNN + features", "both", 1.064, 19, 4, 1.179, "434 KB", 13.0, 1.203, 61, 13),
    ("GRU + features", "both", 1.058, 18, 7, 0.779, "162 KB", 56.9, 1.162, 49, 20),
    ("GRU (RNN)", "TAC sequence", 1.035, 11, 2, 0.147, "94 KB", 134.1, 1.222, 67, 12),
    ("MLP", "features", 1.024, 7, 10, 0.703, "92 KB", 0.7, 1.116, 35, 24),
]

INK, MUTED, BLUE, LIGHT, GREEN, RED, AMBER = (
    "#1f2937", "#6b7280", "#1d4ed8", "#93c5fd", "#15803d", "#b91c1c", "#b45309")

plt.rcParams.update({
    "font.family": "serif", "font.serif": ["Times New Roman", "DejaVu Serif"],
    "font.size": 11, "axes.titlesize": 12, "axes.labelsize": 11,
    "axes.spines.top": False, "axes.spines.right": False,
    "axes.edgecolor": "#9ca3af", "axes.labelcolor": INK, "xtick.color": INK,
    "ytick.color": INK, "figure.dpi": 100, "savefig.dpi": 200,
})


def _geomean(xs):
    xs = [x for x in xs if x > 0]
    return math.exp(sum(math.log(x) for x in xs) / len(xs)) if xs else float("nan")


def _save(fig, name):
    fig.savefig(os.path.join(FIG, name), bbox_inches="tight", facecolor="white")
    plt.close(fig)


# --------------------------------------------------------------------------- #
#  diagrams                                                                   #
# --------------------------------------------------------------------------- #
def fig_pipeline():
    stages = [
        ("MiniC source", ".mc file"),
        ("Front end", "lexer · parser\nAST · semantic checks"),
        ("TAC IR + CFG", "basic blocks, dominators,\nnatural loops"),
        ("27 static features", "19 structural +\n8 opportunity"),
        ("ML recommender", "predict speedup of all\n64 combos, abstain margin"),
        ("Pass manager", "6 passes in canonical\norder, to a fixed point"),
        ("C codegen", "value-semantics C\n(wrapper structs)"),
        ("gcc -O0 + timer", "least of 15 runs,\nexit-code check"),
    ]
    fig, ax = plt.subplots(figsize=(10, 4.4))
    ax.set_xlim(0, 4)
    ax.set_ylim(0, 2)
    ax.axis("off")
    pos = [(i, 1) for i in range(4)] + [(3 - i, 0) for i in range(4)]
    for (title, sub), (cx, cy) in zip(stages, pos):
        x, y = cx + 0.08, cy + 0.18
        hl = title in ("ML recommender", "27 static features")
        ax.add_patch(FancyBboxPatch((x, y), 0.84, 0.64, boxstyle="round,pad=0.02,rounding_size=0.05",
                                    fc="#eff6ff" if hl else "#f9fafb",
                                    ec=BLUE if hl else "#9ca3af", lw=1.4))
        ax.text(x + 0.42, y + 0.47, title, ha="center", va="center", fontsize=11.5,
                fontweight="bold", color=INK)
        ax.text(x + 0.42, y + 0.2, sub, ha="center", va="center", fontsize=9, color=MUTED)
    arrow = dict(arrowstyle="-|>", color=INK, lw=1.2)
    for i in range(3):
        ax.annotate("", xy=(i + 1.08, 1.5), xytext=(i + 0.94, 1.5), arrowprops=arrow)
        ax.annotate("", xy=(i + 0.92, 0.5), xytext=(i + 1.06, 0.5), arrowprops=arrow)
    ax.annotate("", xy=(3.5, 0.84), xytext=(3.5, 1.16), arrowprops=arrow)
    ax.text(2.0, 0.04, "Offline: every program × 64 combos is timed → dataset → model training "
            "(GroupKFold by program)", ha="center", fontsize=9.5, color=MUTED, style="italic")
    _save(fig, "fig_pipeline.png")


def fig_passmask():
    fig, ax = plt.subplots(figsize=(9, 2.3))
    ax.set_xlim(0, 6)
    ax.set_ylim(0, 1.6)
    ax.axis("off")
    names = {"CF": "Constant\nfolding", "DCE": "Dead-code\nelim.", "CSE": "Common\nsubexpr.",
             "LICM": "Loop-invariant\ncode motion", "SR": "Strength\nreduction", "LU": "Loop\nunrolling"}
    for i, (bit, ab) in enumerate(PASS_BITS):
        ax.add_patch(FancyBboxPatch((i + 0.06, 0.55), 0.88, 0.8, boxstyle="round,pad=0.01,rounding_size=0.04",
                                    fc="#eff6ff", ec=BLUE, lw=1.2))
        ax.text(i + 0.5, 1.13, ab, ha="center", fontsize=13, fontweight="bold", color=BLUE)
        ax.text(i + 0.5, 0.78, names[ab], ha="center", va="center", fontsize=8.5, color=INK)
        ax.text(i + 0.5, 0.32, f"bit value {bit}", ha="center", fontsize=9, color=MUTED)
    ax.text(3, 0.02, "combo_id = OR of enabled bits:  0 = no passes,  63 = all six,  "
            "e.g. 21 = CF + CSE + SR", ha="center", fontsize=10, color=INK)
    _save(fig, "fig_passmask.png")


# --------------------------------------------------------------------------- #
#  data charts                                                                #
# --------------------------------------------------------------------------- #
def data_charts():
    s = summarize(DATASET)
    pp = s["per_program"]
    best = [p["best_speedup"] for p in pp]
    allon = [p["all_on_speedup"] for p in pp]

    # 1. distribution of best-combo speedups
    edges = [1.0, 1.05, 1.1, 1.2, 1.35, 1.6, 2.0, 99]
    labels = ["1.00–1.05", "1.05–1.10", "1.10–1.20", "1.20–1.35", "1.35–1.60", "1.60–2.00", "≥ 2.00"]
    hist = [sum(1 for v in best if lo <= v < hi) for lo, hi in zip(edges, edges[1:])]
    fig, ax = plt.subplots(figsize=(8, 3.6))
    bars = ax.bar(labels, hist, color=BLUE, width=0.7)
    for b, v in zip(bars, hist):
        ax.text(b.get_x() + b.get_width() / 2, v + 4, str(v), ha="center", fontsize=10)
    ax.set_xlabel("speedup of the best pass combination over no optimisation")
    ax.set_ylabel("programs")
    ax.set_title(f"Best-combo speedup over {s['n_programs']} programs "
                 f"(geomean ×{s['geomean_best']:.2f}, max ×{s['max_speedup']:.2f})")
    _save(fig, "fig_hist.png")

    # 2. per-category: best vs all-on
    bycat_best, bycat_all = defaultdict(list), defaultdict(list)
    for p in pp:
        bycat_best[p["category"]].append(p["best_speedup"])
        bycat_all[p["category"]].append(p["all_on_speedup"])
    cats = sorted(bycat_best, key=lambda c: -_geomean(bycat_best[c]))
    fig, ax = plt.subplots(figsize=(8, 3.6))
    x = range(len(cats))
    gb = [_geomean(bycat_best[c]) for c in cats]
    ga = [_geomean(bycat_all[c]) for c in cats]
    ax.bar([i - 0.2 for i in x], gb, 0.38, color=BLUE, label="best combo (oracle)")
    ax.bar([i + 0.2 for i in x], ga, 0.38, color=LIGHT, label="all six passes on")
    for i, (a, b) in enumerate(zip(gb, ga)):
        ax.text(i - 0.2, a + 0.03, f"×{a:.2f}", ha="center", fontsize=9)
        ax.text(i + 0.2, b + 0.03, f"×{b:.2f}", ha="center", fontsize=9)
    ax.set_xticks(list(x))
    ax.set_xticklabels([f"{c}\n(n={len(bycat_best[c])})" for c in cats])
    ax.axhline(1.0, color=INK, lw=0.8, ls="--")
    ax.set_ylabel("geomean speedup")
    ax.set_ylim(0.8, max(gb) + 0.35)
    ax.legend(frameon=False)
    ax.set_title("Speedup by benchmark category")
    _save(fig, "fig_category.png")

    # 3. all-on vs best scatter
    fig, ax = plt.subplots(figsize=(6.2, 5))
    ax.scatter(allon, best, s=9, alpha=0.55, color=BLUE, edgecolors="none")
    hi = max(best) + 0.1
    lo = min(min(allon), 0.7)
    ax.plot([lo, hi], [lo, hi], color=INK, lw=0.8, ls="--")
    ax.text(hi - 0.05, hi - 0.25, "best = all-on", ha="right", fontsize=9, color=MUTED, rotation=38)
    ax.axvline(1.0, color=RED, lw=0.7, ls=":")
    ax.set_xlabel("speedup with all six passes on (combo 63)")
    ax.set_ylabel("speedup with the best combo")
    ax.set_title(f"All-on is below the per-program best on "
                 f"{s['n_all_on_suboptimal']} / {s['n_programs']} programs")
    _save(fig, "fig_scatter.png")

    # 4. how often each pass is in the best combo
    freq = [sum(1 for p in pp if p["best_combo"] & bit) / len(pp) * 100 for bit, _ in PASS_BITS]
    fig, ax = plt.subplots(figsize=(7, 3.2))
    bars = ax.bar([ab for _, ab in PASS_BITS], freq, color=BLUE, width=0.6)
    for b, v in zip(bars, freq):
        ax.text(b.get_x() + b.get_width() / 2, v + 1.5, f"{v:.0f}%", ha="center", fontsize=10)
    ax.set_ylim(0, 105)
    ax.set_ylabel("% of programs")
    ax.set_title("How often each pass is part of the program's best combination")
    _save(fig, "fig_passfreq.png")

    # 5. top programs, best vs all-on
    top = sorted(pp, key=lambda p: -p["best_speedup"])[:15][::-1]
    fig, ax = plt.subplots(figsize=(8, 5.2))
    y = range(len(top))
    ax.barh([i + 0.2 for i in y], [p["best_speedup"] for p in top], 0.38, color=BLUE, label="best combo")
    ax.barh([i - 0.2 for i in y], [p["all_on_speedup"] for p in top], 0.38, color=LIGHT, label="all passes on")
    ax.set_yticks(list(y))
    ax.set_yticklabels([f"{p['program_id']}  ({p['category']})" for p in top], fontsize=9)
    for i, p in enumerate(top):
        ax.text(p["best_speedup"] + 0.03, i + 0.2, f"×{p['best_speedup']:.2f}", va="center", fontsize=8.5)
    ax.axvline(1.0, color=INK, lw=0.8, ls="--")
    ax.set_xlabel("speedup over no optimisation")
    ax.legend(loc="lower right", frameon=False)
    ax.set_title("The 15 programs with the largest gains")
    _save(fig, "fig_top.png")

    # 6. recommender vs baselines
    names = ["No passes\n(combo 0)", "All passes\n(combo 63)", "ML recommender\n(held-out)", "Oracle\n(hindsight)"]
    vals = [1.00, 1.19, 1.21, 1.33]
    cols = ["#d1d5db", LIGHT, BLUE, "#1e3a8a"]
    fig, ax = plt.subplots(figsize=(7, 3.4))
    bars = ax.bar(names, vals, color=cols, width=0.6)
    for b, v in zip(bars, vals):
        ax.text(b.get_x() + b.get_width() / 2, v + 0.01, f"×{v:.2f}", ha="center", fontsize=11)
    ax.set_ylim(0.9, 1.4)
    ax.set_ylabel("mean / geomean speedup")
    ax.set_title("Recommender vs fixed policies (900 programs)")
    _save(fig, "fig_reco.png")

    # 7. abstain-margin sweep
    fig, (a1, a2) = plt.subplots(1, 2, figsize=(9, 3.4))
    for (name, tab), c in zip(STRATEGY.items(), [BLUE, AMBER, GREEN]):
        ms = sorted(tab)
        a1.plot(ms, [tab[m][1] for m in ms], "o-", color=c, label=name, ms=4)
        a2.plot(ms, [tab[m][2] for m in ms], "o-", color=c, label=name, ms=4)
    a1.axhline(10, color=RED, lw=0.8, ls="--")
    a1.text(0.18, 10.4, "10% cap", ha="right", fontsize=9, color=RED)
    a1.set_xlabel("abstain margin m (recommend only if predicted > 1+m)")
    a1.set_ylabel("recommendations that regress (%)")
    a2.set_xlabel("abstain margin m")
    a2.set_ylabel("captured oracle gain (%)")
    a2.legend(frameon=False, fontsize=9)
    fig.suptitle("The abstain margin trades a little speedup for far fewer regressions", fontsize=12)
    fig.tight_layout()
    _save(fig, "fig_margin.png")

    # 8. model comparison
    fig, ax = plt.subplots(figsize=(8.5, 4))
    labels = [m[0] for m in MODELS]
    y = range(len(MODELS))[::-1]
    ax.barh([i + 0.2 for i in y], [m[2] for m in MODELS], 0.38, color=BLUE, label="with ≤10% regression cap")
    ax.barh([i - 0.2 for i in y], [m[8] for m in MODELS], 0.38, color=LIGHT, label="uncapped (plain argmax)")
    for i, m in zip(y, MODELS):
        ax.text(m[2] + 0.003, i + 0.2, f"×{m[2]:.3f}", va="center", fontsize=8.5)
        ax.text(m[8] + 0.003, i - 0.2, f"×{m[8]:.3f}  ({m[10]}% regress)", va="center", fontsize=8, color=MUTED)
    ax.set_yticks(list(y))
    ax.set_yticklabels([f"{m[0]}\n[{m[1]}]" for m in MODELS], fontsize=9)
    ax.set_xlim(1.0, 1.29)
    ax.set_xlabel("true speedup of the recommended combo on held-out programs")
    ax.legend(loc="lower right", frameon=False, fontsize=9)
    ax.set_title("Eight-model comparison (same 5-fold program split)")
    _save(fig, "fig_models.png")

    regress_allon = sum(1 for v in allon if v < 0.98)
    print(f"  all-on slower than baseline (>2%) on {regress_allon} / {len(allon)} programs")
    return s


# --------------------------------------------------------------------------- #
#  terminal output -> image                                                   #
# --------------------------------------------------------------------------- #
def _read(name):
    with open(os.path.join(FIG, name), encoding="utf-8-sig") as fh:
        return [l.rstrip("\r\n") for l in fh]


def terminal(name, command, lines, width=92, cols=1, fs=9.6):
    """Render captured command output as a terminal window; cols=2 wraps long listings."""
    if cols == 2:
        cw = (width - 3) // 2
        cut = lambda l: l if len(l) <= cw else l[:cw - 1] + "…"
        half = (len(lines) + 1) // 2
        left, right = lines[:half], lines[half:]
        lines = [f"{cut(a):<{cw}} │ {cut(b)}" for a, b in zip(left, right + [""] * (half - len(right)))]
    lines = [l if len(l) <= width else l[:width - 1] + "…" for l in lines]
    lh = fs * 0.0175
    n = len(lines) + 1
    h = 0.45 + lh * n + 0.1
    fig = plt.figure(figsize=(8.6, h))
    ax = fig.add_axes([0, 0, 1, 1])
    ax.axis("off")
    ax.add_patch(plt.Rectangle((0, 0), 1, 1, color="#0f172a", transform=ax.transAxes))
    ax.add_patch(plt.Rectangle((0, 1 - 0.3 / h), 1, 0.3 / h, color="#1e293b", transform=ax.transAxes))
    for i, c in enumerate(["#ef4444", "#f59e0b", "#22c55e"]):
        ax.plot([0.02 + i * 0.018], [1 - 0.15 / h], "o", ms=6, color=c, transform=ax.transAxes)
    ax.text(0.5, 1 - 0.15 / h, "Windows PowerShell — MiniC optimizer", ha="center", va="center",
            color="#cbd5e1", fontsize=8.5, family="Segoe UI", transform=ax.transAxes)
    body = [("PS> " + command, "#86efac")] + [(l, "#e2e8f0") for l in lines]
    for i, (t, c) in enumerate(body):
        ax.text(0.015, 1 - (0.42 + lh * i) / h, t, color=c, family="Consolas", fontsize=fs,
                va="top", transform=ax.transAxes)
    _save(fig, name)


def terminal_figures():
    terminal("term_tac.png", "python -m src.minic.driver docs/walk.mc --tac", _read("term_tac.txt")[:24], cols=2)
    terminal("term_tac21.png", "python -m src.minic.driver docs/walk.mc --optimize 21 --tac",
             _read("term_tac21.txt")[:27], cols=2)
    terminal("term_cfg.png", "python -m src.minic.driver docs/walk.mc --cfg", _read("term_cfg.txt")[1:29], cols=2)
    terminal("term_features.png", "python -m src.minic.driver docs/walk.mc --features",
             _read("term_features.txt")[:30], cols=2)
    terminal("term_licm.png", "python -m src.minic.driver benchmarks/loops/licm_arith.mc --optimize 9 --tac",
             _read("term_licm.txt")[:38], cols=2)
    un = _read("term_unroll.txt")
    terminal("term_unroll.png", "python -m src.minic.driver docs/walk.mc --optimize 32 --tac",
             un[:17] + ["    ... body repeated ×4 (28 instrs) ..."] + un[37:], cols=2)
    ec = _read("term_emitc.txt")
    i0 = next(i for i, l in enumerate(ec) if l.startswith("int distanceSquared"))
    terminal("term_emitc.png", "python -m src.minic.driver canonical_example.mc --optimize 63 --emit-c",
             ec[2:19] + ["..."] + ec[i0:i0 + 26], cols=2)
    rec = _read("term_recommend.txt")
    j = next(i for i, l in enumerate(rec) if l.startswith("ML-recommended"))
    short = [re.sub(r"Constant Folding \(CF\)", "CF", l) for l in rec[j:]]
    for a, b in [("Dead Code Elimination (DCE)", "DCE"), ("Common Subexpression Elimination (CSE)", "CSE"),
                 ("Loop-Invariant Code Motion (LICM)", "LICM"), ("Strength Reduction (SR)", "SR"),
                 ("Loop Unrolling (LU)", "LU")]:
        short = [l.replace(a, b) for l in short]
    terminal("term_recommend.png", "python -m demo.cli recommend benchmarks/loops/licm_arith.mc --verify",
             rec[:2] + ["static + opportunity features (27):   [19 structural features omitted]"]
             + rec[22:31] + [""] + short, width=100)
    terminal("term_run.png", "python -m src.minic.driver canonical_example.mc --optimize 63 --run",
             [l.replace("Constant Folding (CF), Dead Code Elimination (DCE), Common Subexpression "
                        "Elimination (CSE), Loop-Invariant Code Motion (LICM), Strength Reduction (SR), "
                        "Loop Unrolling (LU)", "CF, DCE, CSE, LICM, SR, LU") for l in _read("term_run.txt")])
    pt = _read("term_pytest.txt")
    keep = [l for l in pt if "PASSED" in l or "FAILED" in l or l.startswith("=")]
    keep = [re.sub(r"\s+\[\s*\d+%\]$", "", l) for l in keep]
    head = [l for l in keep if l.startswith("=")][:1]
    tests = [l for l in keep if not l.startswith("=")]
    tail = [l for l in keep if l.startswith("=")][-1:]
    shown = tests[:30] + ([f"... {len(tests) - 30} more PASSED ..."] if len(tests) > 30 else [])
    terminal("term_pytest.png", "python -m pytest -v tests/test_lexer.py tests/test_parser.py tests/test_sema.py … tests/test_harness_ml.py",
             head + shown + tail, width=100)


def figures():
    os.makedirs(FIG, exist_ok=True)
    fig_pipeline()
    fig_passmask()
    s = data_charts()
    terminal_figures()
    return s


# --------------------------------------------------------------------------- #
#  docx helpers                                                               #
# --------------------------------------------------------------------------- #
import copy                                                    # noqa: E402
import docx                                                    # noqa: E402
from docx.enum.section import WD_SECTION                       # noqa: E402,F401
from docx.enum.table import WD_TABLE_ALIGNMENT                 # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER  # noqa: E402
from docx.oxml import OxmlElement, parse_xml                   # noqa: E402
from docx.oxml.ns import nsdecls, qn                           # noqa: E402
from docx.shared import Cm, Inches, Pt, RGBColor               # noqa: E402

from generate_report_docx import set_cell_background, set_cell_margins  # noqa: E402

TNR = "Times New Roman"
MONO = "Consolas"
TEXT_W = Inches(6.27)          # A4 with 1-inch margins (template page setup)


def _font(run, name=TNR, size=None, bold=None, italic=None, color=None):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.insert(0, rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), name)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    return run


def _fix_widths(t, widths):
    t.autofit = False
    tblPr = t._tbl.tblPr
    for tag in ("w:tblW", "w:tblLayout"):
        old = tblPr.find(qn(tag))
        if old is not None:
            tblPr.remove(old)
    tw = OxmlElement("w:tblW")
    tw.set(qn("w:w"), str(int(sum(widths) * 1440)))
    tw.set(qn("w:type"), "dxa")
    tblPr.append(tw)
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "fixed")
    tblPr.append(lay)
    for col, w in zip(t.columns, widths):
        col.width = Inches(w)
    for row in t.rows:
        for c, w in zip(row.cells, widths):
            c.width = Inches(w)


class Report:
    def __init__(self, template):
        self.d = docx.Document(template)
        self.fig_n = 0
        self.tab_n = 0
        self.figs, self.tabs, self.heads = [], [], []
        self._styles()

    # ---- styles ---------------------------------------------------------- #
    def _styles(self):
        st = self.d.styles
        for name, size, before, after in (("Heading 1", 16, 0, 10), ("Heading 2", 13, 12, 6),
                                          ("Heading 3", 12, 10, 4)):
            s = st[name]
            s.font.name = TNR
            s.font.size = Pt(size)
            s.font.bold = True
            s.font.italic = False
            s.font.color.rgb = RGBColor(0, 0, 0)
            rpr = s.element.get_or_add_rPr()
            rf = rpr.find(qn("w:rFonts"))
            if rf is None:
                rf = OxmlElement("w:rFonts")
                rpr.insert(0, rf)
            for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
                rf.set(qn(a), TNR)
            for a in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
                if rf.get(qn(a)) is not None:
                    del rf.attrib[qn(a)]
            pf = s.paragraph_format
            pf.space_before = Pt(before)
            pf.space_after = Pt(after)
            pf.line_spacing = 1.15
            pf.keep_with_next = True

    # ---- title page ------------------------------------------------------ #
    def title_page(self, title, subtitle, members, month):
        body = self.d.element.body
        els = list(body.iterchildren())
        paras = {i: docx.text.paragraph.Paragraph(e, self.d) for i, e in enumerate(els)
                 if e.tag == qn("w:p")}
        self._set_text(paras[3], title.upper(), size=16, bold=True)
        self._set_text(paras[4], subtitle, size=13, italic=True)
        paras[4].alignment = WD_ALIGN_PARAGRAPH.CENTER
        self._set_text(paras[14], month, size=12, bold=True)
        t = self.d.tables[0]
        reg, name = t.rows[0].cells
        rp = reg.paragraphs
        while len(rp) < len(members) + 1:
            reg._element.append(copy.deepcopy(rp[-1]._element))
            rp = reg.paragraphs
        np_ = name.paragraphs
        while len(np_) < len(members) + 1:
            name._element.append(copy.deepcopy(np_[-1]._element))
            np_ = name.paragraphs
        for i, (r, n) in enumerate(members, start=1):
            self._set_text(rp[i], r, size=14, bold=True)
            self._set_text(np_[i], n, size=14, bold=True)
        for extra in rp[len(members) + 1:]:
            extra._element.getparent().remove(extra._element)
        # drop the template's placeholder headings (everything after the date)
        for e in els[15:]:
            if e.tag != qn("w:sectPr"):
                body.remove(e)

    @staticmethod
    def _set_text(p, text, size=12, bold=False, italic=False):
        for r in list(p.runs):
            r._element.getparent().remove(r._element)
        _font(p.add_run(text), size=size, bold=bold, italic=italic)

    # ---- page furniture -------------------------------------------------- #
    def page_numbers(self):
        sec = self.d.sections[0]
        sec.different_first_page_header_footer = True
        p = sec.footer.paragraphs[0] if sec.footer.paragraphs else sec.footer.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run()
        _font(r, size=10)
        for kind, txt in (("begin", None), (None, "PAGE"), ("separate", None), (None, None), ("end", None)):
            if kind:
                fc = OxmlElement("w:fldChar")
                fc.set(qn("w:fldCharType"), kind)
                r._element.append(fc)
            elif txt:
                it = OxmlElement("w:instrText")
                it.set(qn("xml:space"), "preserve")
                it.text = txt
                r._element.append(it)
            else:
                t = OxmlElement("w:t")
                t.text = "1"
                r._element.append(t)

    def page_break(self):
        p = self.d.add_paragraph()
        p.add_run().add_break(WD_BREAK.PAGE)

    # ---- text ------------------------------------------------------------ #
    def h1(self, text, new_page=True):
        p = self.d.add_paragraph(style="Heading 1")
        p.paragraph_format.page_break_before = new_page
        _font(p.add_run(text), size=16, bold=True)
        self.heads.append((1, text))
        return p

    def h2(self, text):
        p = self.d.add_paragraph(style="Heading 2")
        _font(p.add_run(text), size=13, bold=True)
        self.heads.append((2, text))
        return p

    def h3(self, text):
        p = self.d.add_paragraph(style="Heading 3")
        _font(p.add_run(text), size=12, bold=True)
        return p

    def _runs(self, p, text, size=12):
        for part in re.split(r"(\*\*[^*]+\*\*|`[^`]+`|_[^_]+_)", text):
            if not part:
                continue
            if part.startswith("**"):
                _font(p.add_run(part[2:-2]), size=size, bold=True)
            elif part.startswith("`"):
                _font(p.add_run(part[1:-1]), name=MONO, size=size - 1.5)
            elif part.startswith("_") and part.endswith("_") and len(part) > 2:
                _font(p.add_run(part[1:-1]), size=size, italic=True)
            else:
                _font(p.add_run(part), size=size)

    def para(self, text, align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=6, indent=False):
        p = self.d.add_paragraph()
        pf = p.paragraph_format
        pf.line_spacing = 1.4
        pf.space_after = Pt(after)
        pf.space_before = Pt(0)
        if indent:
            pf.first_line_indent = Cm(0.75)
        p.alignment = align
        self._runs(p, text)
        return p

    def bullets(self, items, numbered=False):
        for i, it in enumerate(items, 1):
            p = self.d.add_paragraph()
            pf = p.paragraph_format
            pf.line_spacing = 1.4
            pf.space_after = Pt(2)
            pf.left_indent = Cm(1.0)
            pf.first_line_indent = Cm(-0.5)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            _font(p.add_run(f"{i}.\t" if numbered else "•\t"), size=12)
            pf.tab_stops.add_tab_stop(Cm(1.0))
            self._runs(p, it)

    def code(self, lines, size=9):
        t = self.d.add_table(rows=1, cols=1)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        _fix_widths(t, [6.27])
        c = t.rows[0].cells[0]
        c.width = TEXT_W
        set_cell_background(c, "F3F4F6")
        set_cell_margins(c, 80, 80, 160, 160)
        first = True
        for ln in lines:
            p = c.paragraphs[0] if first else c.add_paragraph()
            first = False
            pf = p.paragraph_format
            pf.space_after = Pt(0)
            pf.line_spacing = 1.0
            _font(p.add_run(ln if ln else " "), name=MONO, size=size)
        self.d.add_paragraph().paragraph_format.space_after = Pt(2)

    # ---- figures & tables ------------------------------------------------ #
    def figure(self, path, caption, width=6.0):
        self.fig_n += 1
        p = self.d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(2)
        p.add_run().add_picture(path, width=Inches(width))
        c = self.d.add_paragraph()
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        c.paragraph_format.space_after = Pt(10)
        _font(c.add_run(f"Figure {self.fig_n}: "), size=10, bold=True)
        _font(c.add_run(caption), size=10, italic=True)
        self.figs.append((self.fig_n, caption))
        return self.fig_n

    def table(self, header, rows, caption, widths=None, size=10, align=None):
        self.tab_n += 1
        cp = self.d.add_paragraph()
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cp.paragraph_format.keep_with_next = True
        cp.paragraph_format.space_before = Pt(6)
        cp.paragraph_format.space_after = Pt(4)
        _font(cp.add_run(f"Table {self.tab_n}: "), size=10, bold=True)
        _font(cp.add_run(caption), size=10, italic=True)
        self.tabs.append((self.tab_n, caption))
        t = self.d.add_table(rows=1 + len(rows), cols=len(header))
        t.style = self.d.styles["Table Grid"]
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        widths = widths or [6.27 / len(header)] * len(header)
        _fix_widths(t, widths)
        for ri, row in enumerate([header] + rows):
            cells = t.rows[ri].cells
            for ci, val in enumerate(row):
                c = cells[ci]
                c.width = Inches(widths[ci])
                set_cell_margins(c, 40, 40, 80, 80)
                p = c.paragraphs[0]
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                a = (align or ["l"] * len(header))[ci]
                p.alignment = {"l": WD_ALIGN_PARAGRAPH.LEFT, "c": WD_ALIGN_PARAGRAPH.CENTER,
                               "r": WD_ALIGN_PARAGRAPH.RIGHT}[a]
                if ri == 0:
                    set_cell_background(c, "DBEAFE")
                    _font(p.add_run(str(val)), size=size, bold=True)
                else:
                    self._runs(p, str(val), size=size)
            if ri == 0:
                trpr = t.rows[0]._tr.get_or_add_trPr()
                h = OxmlElement("w:tblHeader")
                trpr.append(h)
        self.d.add_paragraph().paragraph_format.space_after = Pt(2)
        return self.tab_n

    # ---- contents pages (static, page numbers from a rendered pass) ------- #
    def contents(self, pages=None):
        pages = pages or {}

        def line(text, level, key, size=12, bold=False):
            p = self.d.add_paragraph()
            pf = p.paragraph_format
            pf.space_after = Pt(0 if level > 1 else 2)
            pf.line_spacing = 1.05
            pf.left_indent = Cm(0.8 * (level - 1))
            pf.tab_stops.add_tab_stop(TEXT_W, WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
            _font(p.add_run(f"{text}\t{pages.get(key, '00')}"), size=size, bold=bold)

        hp = self.d.add_paragraph()
        _font(hp.add_run("TABLE OF CONTENTS"), size=16, bold=True)
        hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        hp.paragraph_format.space_after = Pt(12)
        for level, text in CONTENTS_HEADINGS:
            line(text, level, ("h", text), size=12 if level == 1 else 11, bold=level == 1)
        self.page_break()
        hp = self.d.add_paragraph()
        _font(hp.add_run("LIST OF FIGURES"), size=16, bold=True)
        hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        hp.paragraph_format.space_after = Pt(12)
        for n, cap in FIGURE_LIST:
            line(f"Figure {n}: {cap}", 2, ("f", n), size=10.5)
        hp = self.d.add_paragraph()
        hp.paragraph_format.space_before = Pt(18)
        _font(hp.add_run("LIST OF TABLES"), size=16, bold=True)
        hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        hp.paragraph_format.space_after = Pt(12)
        for n, cap in TABLE_LIST:
            line(f"Table {n}: {cap}", 2, ("t", n), size=10.5)

    def save(self, path):
        self.d.core_properties.title = "AI-Based Compiler Optimization for MiniC"
        self.d.core_properties.author = "Ashmit Verma, S Karthik Sudharsan, Mohammed Nadeem"
        self.d.save(path)


# filled in by a first (dry) build so the contents pages can list everything
CONTENTS_HEADINGS, FIGURE_LIST, TABLE_LIST = [], [], []


def F(name):
    return os.path.join(FIG, name)


def S(name):
    return os.path.join(ROOT, "output-images", name)


# --------------------------------------------------------------------------- #
#  the report                                                                 #
# --------------------------------------------------------------------------- #
def write_body(R):
    # ===================================================================== 1
    R.h1("1. Abstract")
    R.para(
        "Production compilers such as GCC and LLVM expose optimisation through a handful of fixed "
        "levels (`-O1`, `-O2`, `-O3`), each of which applies the same pass sequence to every program. "
        "Optimisation passes, however, interact: one pass can create or destroy the opportunities of "
        "another, and a pass that speeds up one program can slow another down. Choosing _which_ "
        "passes to run for a particular program is therefore a per-program decision that a fixed "
        "level cannot make.", indent=True)
    R.para(
        "This project builds a complete optimising compiler for **MiniC**, a small, statically typed, "
        "pointer-free subset of C, and a machine-learning recommender that decides which of its six "
        "optimisation passes to enable for a program it has never executed. The compiler consists of a "
        "hand-written lexer and recursive-descent parser, a semantic analyser, lowering to three-address "
        "code (TAC), control-flow-graph and loop analysis, six independently switchable passes (constant "
        "folding, dead-code elimination, common-subexpression elimination, loop-invariant code motion, "
        "strength reduction and loop unrolling, giving 2⁶ = 64 combinations) and a C code generator that "
        "preserves MiniC's value semantics. Because MiniC has no pointers, there is no aliasing, which "
        "makes every pass provably sound without alias analysis.", indent=True)
    R.para(
        "Every result in this report is measured rather than modelled. Each of **900 benchmark programs** "
        "(30 hand-written and 870 generated) was compiled at all 64 pass combinations, built with "
        "`gcc -O0` so that the C compiler contributes no optimisation of its own, and timed on real "
        "hardware: **57,600 timed configurations**, every one verified to return the same result as the "
        "unoptimised program. Choosing the best combination per program gives a geometric-mean speedup of "
        "**×1.31** (up to ×4.37), while simply enabling all six passes gives ×1.19 and is not the best "
        "choice for 715 of the 900 programs. No fixed non-empty combination avoids slowing at least one "
        "program down.", indent=True)
    R.para(
        "From 27 static features of a program's TAC, a Random Forest recommender with a risk-aware "
        "_abstain margin_ achieves **×1.21** on held-out programs, capturing **63%** of the speedup an "
        "oracle could obtain, while only **9%** of its recommendations make the program slower. An "
        "eight-model comparison (linear, Random Forest, Histogram Gradient Boosting, MLP, 1D-CNN and GRU "
        "over raw TAC, and hybrids) shows that tree ensembles over engineered features perform best, and "
        "that a CNN reading only the raw instruction sequence comes within 1% of them.", indent=True)
    R.para("**Keywords:** compiler optimisation, pass selection, phase ordering, three-address code, "
           "machine learning, gradient boosting, random forest, MiniC.", after=0)

    # ===================================================================== 2
    R.h1("2. Introduction")
    R.h2("2.1 Background")
    R.para(
        "An optimising compiler translates a program into an intermediate representation (IR) and applies "
        "a sequence of _passes_, each of which rewrites the IR into an equivalent but cheaper form. Classic "
        "passes include constant folding, which evaluates expressions whose operands are known at compile "
        "time; dead-code elimination, which removes computations whose results are never used; "
        "common-subexpression elimination, which computes a repeated expression only once; loop-invariant "
        "code motion, which hoists computations that do not change inside a loop to just before it; "
        "strength reduction, which replaces an expensive operation such as a multiplication by a cheaper "
        "addition; and loop unrolling, which replicates a loop body to reduce branch overhead [1], [2], [12].",
        indent=True)
    R.para(
        "Real compilers ship dozens of such passes and expose them to users as a few optimisation levels. "
        "The levels are chosen by compiler engineers to perform well _on average_ over a benchmark suite, "
        "and every program compiled at a given level receives the same passes in the same order.",
        indent=True)
    R.h2("2.2 Problem Statement")
    R.para(
        "Two related problems make a fixed level sub-optimal. The **pass-selection problem** asks which "
        "subset of passes to apply to a given program; the **phase-ordering problem** asks in which order "
        "to apply them [7]. Passes interact in both directions. Constant folding can turn `i * scale` into "
        "`i * 8`, which is exactly the form strength reduction needs; on the other hand, unrolling a small "
        "loop increases code size and can make it slower, and strength reduction can replace a "
        "multiplication that was already cheap by additional loop-carried work. The best subset is "
        "therefore program-dependent.", indent=True)
    R.para(
        "The goal of this project is to answer, for an unseen MiniC program, the question: _which of the "
        "six available optimisation passes should be enabled?_ The answer must be produced statically, "
        "without running the program, and must be measurably better than both doing nothing and enabling "
        "everything, without making programs slower.", indent=True)
    R.h2("2.3 Motivation")
    R.para(
        "Searching all combinations of passes for every program (iterative compilation) finds the best "
        "configuration but requires compiling and running the program many times, which is impractical in "
        "a normal build. Machine learning offers a middle path: learn, from programs that _have_ been "
        "measured exhaustively, a mapping from static program properties to the expected benefit of each "
        "configuration, and apply it to new programs at compile time [3], [7]. The project studies "
        "this idea end-to-end in a setting small enough to measure every configuration exactly, so that "
        "the recommender can be compared against the true optimum rather than an estimate.", indent=True)
    R.h2("2.4 Objectives")
    R.bullets([
        "Design and implement a complete compiler front end, IR and C back end for MiniC.",
        "Implement six independently switchable, provably sound optimisation passes and a pass manager "
        "that iterates them to a fixed point.",
        "Extract a static feature vector from each program that describes both its structure and its "
        "optimisation opportunities.",
        "Build a benchmark corpus and a timing harness that measures every program at every pass "
        "combination, with a correctness check on every build.",
        "Train and validate a recommender that selects a pass combination for unseen programs while "
        "keeping the rate of slowdowns under 10%.",
        "Compare classical and neural models, including models that read the raw TAC instruction "
        "sequence instead of engineered features.",
        "Provide a command-line tool and a browser playground that demonstrate the full pipeline.",
    ], numbered=True)
    R.h2("2.5 Scope and Assumptions")
    R.para(
        "The optimiser works at the TAC level and emits C; the C is compiled with `gcc -O0`. Pinning the C "
        "compiler at `-O0` is deliberate: at that level GCC performs essentially no optimisation, so the "
        "difference between two builds is attributable entirely to this project's passes. `gcc -O2` is "
        "reported only as a reference point. The passes are intra-procedural; inter-procedural "
        "optimisation and inlining are outside the scope of this work. All measurements were taken on a "
        "single Windows 11 machine.", indent=True)
    R.h2("2.6 Team and Contributions")
    R.table(["Member", "Reg. No.", "Area of contribution"], [
        ["Ashmit Verma", "24BCT0301", "Front end (lexer, recursive-descent parser, AST, semantic analysis), "
                                      "TAC lowering, CFG / dominator / loop analysis, 19 structural features."],
        ["S Karthik Sudharsan", "24BAI0243", "Six optimisation passes, 64-combination pass manager with fixed-point "
                                             "iteration, value-semantics C code generator."],
        ["Mohammed Nadeem", "23BCE2091", "Timing harness and correctness oracle, benchmark corpus and generator, "
                                         "opportunity features, ML recommender and model comparison, "
                                         "CLI and web playground."],
    ], "Team members and areas of contribution", widths=[1.6, 1.0, 3.67])
    R.h2("2.7 Organisation of the Report")
    R.para(
        "Section 3 describes what is novel in the approach. Section 4 presents the proposed work: the "
        "architecture, the MiniC language, every compiler stage, the measurement method, the corpus and "
        "the learning formulation. Section 5 covers the implementation, screenshots of the working system "
        "and the experimental results. Section 6 concludes with limitations and future work, and Section 7 "
        "gives the repository link.", indent=True)

    # ===================================================================== 3
    R.h1("3. Novelty in Project", new_page=False)
    R.para(
        "Machine-learning-guided compilation has been studied for over a decade, from MILEPOST GCC, which "
        "predicted good optimisation flags from static program features [3], to reinforcement-learning "
        "environments for phase ordering such as CompilerGym and AutoPhase [4], [5] and to production use "
        "in LLVM's MLGO framework for inlining and register-allocation decisions [6]. This project does not "
        "claim a new learning algorithm. Its contribution is a complete, small, fully measured system in "
        "which every design choice serves the trustworthiness of the final number. The following aspects "
        "distinguish it.", indent=True)
    R.h2("3.1 An Alias-Free Language Designed for Sound Optimisation")
    R.para(
        "MiniC was designed around a single invariant: **no two variables can ever refer to the same memory**. "
        "It excludes pointers, dynamic allocation, unions and function pointers, and gives arrays, structs "
        "and strings value semantics (assignment and parameter passing copy them). Alias analysis is one "
        "of the hardest problems in real optimisers; removing aliasing at the language level makes every "
        "pass in this project provably sound with straightforward data-flow reasoning, so that the measured "
        "speedups can never come from an incorrect transformation.", indent=True)
    R.h2("3.2 Measured, Not Modelled — with a Correctness Gate")
    R.para(
        "The dataset is not produced by a cost model or a simulator. Every program is built and run at all "
        "64 configurations, and a timing is recorded **only** if the optimised binary returns exactly the "
        "same value as the unoptimised one. Across 57,600 builds, no configuration failed this check. The "
        "ground truth is therefore an exhaustive, verified search, which allows the recommender to be "
        "scored against the true per-program optimum (the oracle).", indent=True)
    R.h2("3.3 Opportunity Features")
    R.para(
        "Besides 19 conventional structural features (instruction counts, loop depth, branch density and "
        "so on), the compiler computes eight **opportunity features**: cheap static estimates of what each "
        "pass would actually change, such as the number of foldable constant expressions, hoistable "
        "loop-invariant instructions or redundant subexpressions, and an estimate of the dynamic "
        "instruction count. They give the model a preview of each pass's pay-off without running anything.",
        indent=True)
    R.h2("3.4 Risk-Aware Recommendation")
    R.para(
        "A regressor that simply picks the configuration with the highest predicted speedup makes "
        "confident mistakes: on this corpus, plain argmax slows down 12–14% of programs. The recommender "
        "therefore **abstains** — recommends no passes — unless the best predicted speedup exceeds a margin. "
        "Crucially, the margin, the model family and the selection rule (argmax, margin or per-pass "
        "classifiers) are themselves chosen by cross-validation, as the most useful strategy whose "
        "held-out regression rate stays under 10%.", indent=True)
    R.h2("3.5 Leak-Free Evaluation")
    R.para(
        "A program contributes 64 rows to the dataset, all with identical static features. A random "
        "row-level split would place rows of the same program in both training and test sets and grossly "
        "overstate accuracy. All evaluation in this project uses **GroupKFold by program**: an entire "
        "program is held out, so every reported number describes programs the model has never seen.",
        indent=True)
    R.h2("3.6 Learning Directly from the IR")
    R.para(
        "In addition to feature-based models, the project tokenises each program's TAC into an "
        "opcode/operand-kind sequence and trains 1D convolutional and recurrent (GRU) networks on it. All "
        "eight models are scored on one shared split with the same recommendation rule, which makes the "
        "comparison between hand-engineered features and learned representations direct.", indent=True)

    # ===================================================================== 4
    R.h1("4. Proposed Work")
    R.h2("4.1 System Architecture")
    R.para(
        "Figure 1 shows the end-to-end pipeline. A MiniC source file passes through the front end into "
        "three-address code, from which the control-flow graph and the 27 features are computed. At compile "
        "time the recommender scores all 64 pass combinations from those features and selects one; the "
        "pass manager applies it; the code generator emits C, which is compiled with `gcc -O0`. Offline, the "
        "same pipeline runs every program at every combination to produce the training data.", indent=True)
    R.figure(F("fig_pipeline.png"), "End-to-end architecture of the MiniC optimiser and recommender", 6.2)
    R.h2("4.2 The MiniC Language")
    R.para(
        "MiniC supports `int`, `float` and `char`; fixed-size one- and two-dimensional arrays; non-recursive "
        "structs (which may contain arrays and other structs); fixed-size strings (`char` arrays with "
        "compile-time length checks); functions and recursion; `if`/`else`, `while` and `for`; block scoping; "
        "and the usual arithmetic, relational and logical operators. Pointers, dynamic memory, `goto`, "
        "`switch`, unions, variadic functions and the preprocessor are excluded. A program has no I/O: its "
        "observable result is the value returned by `main()`. An abridged grammar is shown below.",
        indent=True)
    R.code([
        "program     → (struct_decl | func_decl | var_decl)*",
        "struct_decl → 'struct' IDENT '{' field_decl+ '}' ';'",
        "var_decl    → type_spec IDENT array_suffix? ('=' expr)? ';'",
        "array_suffix→ '[' INT_LIT ']' ('[' INT_LIT ']')?",
        "func_decl   → type_spec IDENT '(' param_list? ')' block",
        "stmt        → var_decl | expr_stmt | if_stmt | while_stmt | for_stmt | return_stmt | block",
        "expr        → assignment ;  assignment → postfix '=' assignment | logic_or",
        "logic_or    → logic_and ('||' logic_and)* ;  logic_and → equality ('&&' equality)*",
        "additive    → multiplicative (('+'|'-') multiplicative)*",
        "unary       → ('-'|'!') unary | postfix",
    ])
    R.h2("4.3 Front End")
    R.para(
        "The **lexer** (`frontend/lexer.py`) converts source text into tokens tagged with kind and line "
        "number. The **parser** (`frontend/parser.py`) is a hand-written recursive-descent parser with one "
        "function per grammar rule, producing an abstract syntax tree of typed nodes (`Program`, `FuncDecl`, "
        "`StructDecl`, `WhileStmt`, `BinaryExpr`, `ArrayAccessExpr` and others). The **semantic analyser** "
        "(`frontend/sema.py`) then checks declarations before use, function arity, operand types, integer "
        "array subscripts, valid l-values, non-recursive struct sizes, missing returns and duplicate "
        "declarations, reporting errors with line, column and a caret under the offending token.", indent=True)
    R.h2("4.4 Three-Address Code and Control-Flow Analysis")
    R.para(
        "The IR generator lowers the AST into TAC: a flat list of instructions with at most one operator "
        "each, explicit temporaries and explicit control flow through labels, `goto` and `ifFalse`. Figure 2 "
        "shows the TAC of the running example `docs/walk.mc`, a loop that accumulates `i * scale` twice per "
        "iteration. Nested aggregate assignments such as `a[i].f = v` are lowered to a load / modify-copy / "
        "store-back chain so that value semantics are preserved.", indent=True)
    R.figure(F("term_tac.png"), "Three-address code generated for the running example (no optimisation)", 5.6)
    R.para(
        "The CFG builder (`ir/cfg.py`) partitions the TAC into basic blocks, connects them with edges, "
        "computes dominators and identifies natural loops from back edges (Figure 3). Loop-invariant code "
        "motion, strength reduction and unrolling all depend on this loop structure.", indent=True)
    R.figure(F("term_cfg.png"), "Basic blocks and loop structure of the running example", 5.6)
    R.h2("4.5 Static Feature Extraction")
    R.para(
        "For every program the compiler computes a 27-dimensional feature vector without executing it "
        "(Table 2). The 19 structural features describe size, control flow and the instruction mix; the 8 "
        "opportunity features estimate how much each pass could change the program. Figure 4 shows the "
        "vector for the running example; note `opp_cse_redundant = 1`, the repeated `i * scale`.",
        indent=True)
    R.table(["Group", "Features"], [
        ["Size & control flow", "total_instructions, basic_block_count, loop_count, max_loop_depth, "
                                "branch_count, branch_density, cyclomatic_complexity"],
        ["Instruction mix", "arithmetic_ops_count, multiplication_count, constant_load_count, "
                            "function_call_count, recursive_call_count, string_ops_count"],
        ["Data access", "array_access_count, array_2d_access_count, struct_access_count, "
                        "named_variable_count, temp_variable_count, instruction_density_in_loops"],
        ["Opportunity (8)", "opp_const_foldable, opp_cse_redundant, opp_licm_hoistable, opp_sr_reducible, "
                            "opp_unrollable_loops, opp_loop_body_insts, opp_hot_invariant_frac, "
                            "opp_est_dyn_kilo_insts"],
    ], "The 27 static features (19 structural + 8 opportunity)", widths=[1.6, 4.67])
    R.figure(F("term_features.png"), "Feature vector extracted for the running example", 5.6)
    R.h2("4.6 Optimisation Passes")
    R.para(
        "Each pass is a function from TAC to TAC and is controlled by one bit of a 6-bit mask (Figure 5); a "
        "combination is identified by the integer value of the mask, from 0 (no passes) to 63 (all six). "
        "Table 3 summarises the passes.", indent=True)
    R.figure(F("fig_passmask.png"), "The six passes and their bits in the combination mask", 6.0)
    R.table(["Bit", "Pass", "Transformation"], [
        ["1", "Constant folding (CF)", "Evaluates constant expressions with 32-bit C semantics, propagates "
                                       "single-assignment constants across the function, applies algebraic "
                                       "identities and removes constant branches."],
        ["2", "Dead-code elimination (DCE)", "Removes instructions whose results are never used, unreachable "
                                             "blocks and unused allocations (backward liveness)."],
        ["4", "Common-subexpression elim. (CSE)", "Local value numbering with copy propagation: a repeated "
                                                   "expression is computed once and reused."],
        ["8", "Loop-invariant code motion (LICM)", "Hoists instructions whose operands do not change in the "
                                                    "loop into a labelled pre-header."],
        ["16", "Strength reduction (SR)", "Replaces `i * k` for an induction variable `i` with an "
                                          "accumulator updated by `+ k` each iteration."],
        ["32", "Loop unrolling (LU)", "Replicates a counted loop body four times and keeps the original loop "
                                      "as a remainder loop."],
    ], "Summary of the optimisation passes", widths=[0.45, 1.85, 3.97], align=["c", "l", "l"])
    R.para(
        "Passes enable each other. On the running example, strength reduction alone does nothing because "
        "the multiplier `scale` is a variable; after constant folding has propagated `scale = 8`, SR can "
        "replace both multiplications with a single accumulator `__sr7` that grows by 8 each iteration, and "
        "CSE makes both uses share it (Figure 6).", indent=True)
    R.figure(F("term_tac21.png"), "Combination 21 (CF + CSE + SR): both multiplications replaced by an "
                                  "accumulator", 5.6)
    R.para(
        "Figure 7 shows LICM and CF on the benchmark `licm_arith`, whose loop body consists entirely of "
        "invariant products of `a`, `b` and `c`. All six products are folded to constants and hoisted out of "
        "the loop, leaving only additions inside it. Figure 8 shows loop unrolling: the body is replicated "
        "four times under a guard `i < 997`, followed by the original loop for the remaining iterations.",
        indent=True)
    R.figure(F("term_licm.png"), "Combination 9 (CF + LICM) on licm_arith: invariant products folded "
                                 "and hoisted", 5.6)
    R.figure(F("term_unroll.png"), "Combination 32 (LU): four-way unrolled loop with remainder loop", 5.6)
    R.h2("4.7 Pass Manager")
    R.para(
        "The pass manager (`optimizer/pass_manager.py`) applies the enabled passes in the canonical order "
        "**CF → CSE → LU → LICM → SR → DCE** and repeats the whole sequence until the program stops changing "
        "(a fixed point). Iteration lets each pass clean up after the others: for example, DCE removes the "
        "temporaries that CSE and SR leave unused, and CF can fold expressions that LICM has just hoisted. "
        "The order is fixed so that the 64 combinations are a clean pass-_selection_ problem.", indent=True)
    R.h2("4.8 Value-Semantics Code Generation")
    R.para(
        "The C emitter (`codegen/c_emitter.py`) produces C99 that any GCC can compile. To keep MiniC's copy "
        "semantics, arrays and strings are wrapped in one-field structs (`arr_int_3_3`, `str_6`), so that C "
        "assignment and parameter passing copy them instead of decaying to pointers; struct typedefs are "
        "emitted in dependency order. Figure 9 shows part of the output for `canonical_example.mc`, the "
        "program that exercises every language feature.", indent=True)
    R.figure(F("term_emitc.png"), "Generated value-semantics C for canonical_example.mc (combo 63, excerpt)", 5.6)
    R.h2("4.9 Measurement Methodology")
    R.para(
        "The harness (`harness/`) compiles each emitted C file with `gcc -O0 -std=c99` and runs the binary "
        "after four warm-up runs, keeping the **minimum of 15** wall-clock runs. The sweep runs eight "
        "programs in parallel; the minimum is the least-interfered sample under that load. The speedup of "
        "combination _c_ for program _p_ is defined as",
        indent=True)
    R.code(["speedup(p, c) = time(p, combo 0) / time(p, c)          (> 1 means faster)"], size=10)
    R.para(
        "Before a timing is recorded, the binary's exit code must equal that of combination 0 (the "
        "correctness gate). The full sweep of 900 programs took 3 hours 20 minutes on eight workers.",
        indent=True)
    R.h2("4.10 Benchmark Corpus")
    R.para(
        "The corpus contains 30 hand-written programs in five categories and 870 programs produced by a "
        "parametric generator (`harness/gen_corpus.py`). The generator instantiates single-pattern programs "
        "that target one pass each, with randomised constants, trip counts and loop forms, and weighted "
        "**composites** that combine several patterns in one program. A manifest records each program's "
        "pattern so that results can be broken down by program type (Table 4).", indent=True)
    R.table(["Source", "Category / pattern", "Programs", "Purpose"], [
        ["Hand-written", "loops", "15", "LICM, CSE, constant propagation, polynomial evaluation"],
        ["Hand-written", "numeric", "6", "array kernels: convolution, dot product, histogram, matrix multiply, stencil, vector scale"],
        ["Hand-written", "structs", "4", "bounding boxes, complex arithmetic, distance accumulation, particle simulation"],
        ["Hand-written", "recursion", "3", "Fibonacci, Ackermann, Takeuchi"],
        ["Hand-written", "strings", "2", "Caesar cipher, vowel counting"],
        ["Generated", "cf, cse, licm, sr, div, struct, unroll", "32 each", "one optimisation opportunity per program"],
        ["Generated", "branch, nested, recur, neutral", "31 each", "branchy code, nested loops, recursion, "
                                                               "no opportunity (control)"],
        ["Generated", "composite", "522", "2–4 patterns combined in one program"],
        ["Total", "", "900", "57,600 (program, combination) measurements"],
    ], "Composition of the 900-program benchmark corpus", widths=[1.05, 2.0, 0.8, 2.42],
        align=["l", "l", "c", "l"])
    R.h2("4.11 Learning Formulation")
    R.para(
        "Each (program, combination) pair is one training row: the 27 features of the program, the 6 bits "
        "of the combination, and the measured speedup as the regression target. A regressor _f_ therefore "
        "predicts the speedup of any combination for any program. For a new program the recommender "
        "evaluates all 64 combinations and applies the selection rule",
        indent=True)
    R.code([
        "c* = argmax_c f(x_p, c)",
        "recommend c*      if f(x_p, c*) > 1 + m",
        "recommend combo 0 (no passes) otherwise          m = abstain margin",
    ], size=10)
    R.para(
        "Three regressors were considered for the production recommender: Random Forest (400 trees), Extra "
        "Trees (500 trees) and Histogram Gradient Boosting (600 iterations, learning rate 0.05) [8], [10], "
        "[11]. Each was combined with plain argmax, with margins m ∈ {0.03, 0.05, 0.08, 0.12, 0.18}, and a "
        "per-pass classifier alternative. Every strategy was scored by 5-fold GroupKFold on: the mean true "
        "speedup of the recommended combination; the fraction of the oracle's gain it captures; and the "
        "fraction of programs it makes more than 2% slower. The winner is the strategy with the highest "
        "useful gain whose regression rate is at most 10%.", indent=True)
    R.h2("4.12 Model Comparison Design")
    R.para(
        "To test whether learned program representations can replace engineered features, eight models were "
        "compared under identical conditions (Table 5). Sequence models read each program's unoptimised TAC "
        "converted to tokens (opcode plus operand kind, with names and values dropped; median length 284, "
        "capped at 640). Neural models are trained program-batched: a program is encoded once and its "
        "encoding is combined with each of the 64 pass vectors by a small prediction head. Every model is "
        "evaluated on the same five folds, using the same margin sweep and the same 10% regression cap.",
        indent=True)
    R.table(["Model", "Input", "Architecture"], [
        ["Linear (Ridge)", "features", "standardised features + pass bits, L2-regularised linear regression"],
        ["Random Forest", "features", "400 trees, min. 2 samples per leaf"],
        ["HistGradient Boosting", "features", "600 boosting iterations, 31 leaves, learning rate 0.05"],
        ["MLP", "features", "27 → 128 → H encoder, ReLU, dropout 0.1"],
        ["1D-CNN", "TAC sequence", "token embedding → 3 × Conv1d (kernel 5; 64, 64, 128 channels) → pooling"],
        ["GRU", "TAC sequence", "token embedding → 1-layer GRU, sequence capped at 320 tokens"],
        ["1D-CNN + features", "both", "CNN encoding concatenated with a 64-unit feature encoder"],
        ["GRU + features", "both", "GRU encoding concatenated with a 64-unit feature encoder"],
    ], "Models in the comparison (neural models share a 6-bit pass input and a 64-unit prediction head)",
        widths=[1.6, 1.15, 3.52])

    # ===================================================================== 5
    R.h1("5. Implementation & Screenshots")
    R.h2("5.1 Technology Stack and Environment")
    R.table(["Component", "Technology"], [
        ["Compiler, harness, ML", "Python 3.13.15 (≈ 9,300 lines including tests)"],
        ["C compiler", "GCC 16.2.0 (WinLibs MinGW-w64, UCRT), always at -O0 for measurement"],
        ["Classical ML", "scikit-learn 1.9.1, NumPy 2.5.4, pandas 3.0.6"],
        ["Neural models", "PyTorch 2.14.1 (CPU)"],
        ["Web playground", "Flask 3.1.3 with a single-page HTML/JavaScript front end"],
        ["Plots & report", "Matplotlib 3.11.2, python-docx"],
        ["Testing", "pytest 9.1.1"],
        ["Platform", "Windows 11, 16 logical CPUs, sweep on 8 parallel workers"],
    ], "Implementation technology and experimental environment", widths=[1.8, 4.47])
    R.h2("5.2 Repository Structure")
    R.code([
        "src/minic/",
        "  frontend/   lexer.py  parser.py  ast_nodes.py  sema.py",
        "  ir/         tac.py  ir_generator.py  cfg.py",
        "  features/   extractor.py (19 structural)  opportunity.py (8 opportunity)",
        "  optimizer/  constant_folding.py  dce.py  cse.py  licm.py",
        "              strength_reduction.py  loop_unroll.py  pass_manager.py",
        "  codegen/    c_emitter.py",
        "  harness/    compiler.py  timer.py  sweeper.py  gen_corpus.py  report.py",
        "  ml/         dataset.py  model.py  train.py  predictor.py",
        "              sequences.py  models_torch.py  models_compare.py",
        "  driver.py   command-line compiler",
        "demo/         cli.py (recommend | benchmark)   app.py (browser playground)",
        "benchmarks/   30 hand-written programs + generated/ (870 programs)",
        "tests/        pytest suite: front end, IR, every pass, 64-combo sweeps, harness, ML",
        "run_experiment.py   sweep → dataset → strategy bake-off → train → RESULTS.md",
    ])
    R.h2("5.3 Running the System")
    R.code([
        "pip install -r requirements.txt",
        "python -m src.minic.driver prog.mc --tac --features      # inspect TAC and features",
        "python -m src.minic.driver prog.mc --optimize 63 --run   # optimise, compile, run",
        "python -m demo.cli recommend prog.mc --verify            # ML pick + measured speedup",
        "python -m demo.app                                       # playground, port 5005",
        "python run_experiment.py --compare-models                # full experiment",
    ])
    R.h2("5.4 Screenshots")
    R.para(
        "The browser playground (`demo/app.py`) runs the whole pipeline interactively: the user edits a "
        "MiniC program, the server compiles it unoptimised and optimised, times both builds, checks that "
        "they return the same value and shows the speedup, a per-pass report and a side-by-side TAC diff. "
        "Figures 10–12 show the struct-distance example.", indent=True)
    R.figure(S("run1/image_370.png"), "Playground, TAC diff tab: struct-distance program before (combo 0) "
                                      "and after (combo 63)", 5.3)
    R.figure(S("run1/image_371.png"), "Playground, Optimization tab: passes applied, per-pass effect and "
                                      "the static features read by the model", 5.3)
    R.figure(S("run1/image_372.png"), "Playground, Result tab: ×2.52 speedup (190.57 ms → 75.49 ms); "
                                      "both builds return 130", 5.3)
    R.para(
        "Figures 13–15 show the loop-invariant arithmetic example. Unrolling and LICM replace the six "
        "multiplications in the loop body with additions of folded constants, and the optimised `-O0` build "
        "runs ×1.93 faster than the unoptimised one — faster, in this case, than GCC's own `-O2` build of "
        "the unoptimised C.", indent=True)
    R.figure(S("run2/image_373.png"), "Playground, TAC diff tab: loop-invariant arithmetic program", 5.3)
    R.figure(S("run2/image_374.png"), "Playground, Optimization tab for the loop-invariant program", 5.3)
    R.figure(S("run2/image_375.png"), "Playground, Result tab: ×1.93 speedup (200.64 ms → 104.10 ms)", 5.3)
    R.para(
        "The command-line tool reports the model's recommendation for a program together with its top five "
        "predicted combinations, and with `--verify` it builds and times the program to confirm the "
        "prediction (Figure 16). Figure 17 shows the canonical example, which uses every MiniC feature, "
        "returning its expected value 83 with all six passes enabled.", indent=True)
    R.figure(F("term_recommend.png"), "Recommendation for licm_arith: predicted ×1.68, measured ×1.70 "
                                      "(168.92 ms → 99.62 ms)", 5.6)
    R.figure(F("term_run.png"), "canonical_example.mc optimised with all passes still returns 83", 5.2)

    # ---- results
    R.h2("5.5 Experimental Results")
    R.h3("5.5.1 Overview of the sweep")
    R.table(["Measure", "Value"], [
        ["Programs × combinations", "900 × 64 = 57,600 timed configurations"],
        ["Configurations failing the correctness gate", "0"],
        ["Best combination per program (oracle), geometric mean", "**×1.31**   (maximum ×4.37, complex_arith)"],
        ["All six passes on (combo 63), geometric mean", "×1.19"],
        ["Programs where all-on is not the best (>2% behind)", "715 / 900"],
        ["Programs made >2% slower by all-on", "112 / 900"],
        ["Fixed non-empty combination that never regresses >3%", "none — only combo 0 is safe"],
    ], "Headline results of the 900-program sweep", widths=[3.6, 2.67])
    R.para(
        "The optimiser helps almost every program, but by very different amounts (Figure 18). About 39% of "
        "programs gain less than ×1.2 from even their best combination, while 127 gain more than ×1.6.",
        indent=True)
    R.figure(F("fig_hist.png"), "Distribution of the best-combination speedup over 900 programs", 5.0)
    R.h3("5.5.2 Results by category")
    R.para(
        "Struct-heavy programs benefit most (×2.70 geomean): value semantics make the unoptimised code copy "
        "and re-load struct fields repeatedly, and CSE, LICM and DCE remove most of that work. Loop-dominated "
        "programs follow at ×1.75. Recursive programs gain least (×1.13), because their cost lies in calls, "
        "which the intra-procedural passes do not touch (Figure 19).", indent=True)
    R.figure(F("fig_category.png"), "Geometric-mean speedup by category, best combination vs all passes", 5.3)
    R.h3("5.5.3 No single combination is best")
    R.para(
        "Figure 20 compares, for every program, the speedup of enabling everything with the speedup of the "
        "best combination. Points on the diagonal are programs for which all-on is optimal; most programs "
        "lie well above it, and 112 lie left of the dashed line, where all-on is slower than doing nothing. "
        "Figure 21 shows that no pass belongs to every program's best combination, and Figure 22 lists the "
        "programs with the largest gains.", indent=True)
    R.figure(F("fig_scatter.png"), "All-on speedup vs best-combination speedup, one point per program", 4.6)
    R.figure(F("fig_passfreq.png"), "Share of programs whose best combination includes each pass", 4.0)
    R.figure(F("fig_top.png"), "The fifteen programs with the largest speedups", 5.2)
    R.h3("5.5.4 The recommender")
    R.para(
        "The strategy bake-off selected a **Random Forest with an abstain margin of 0.05**: a combination is "
        "recommended only if its predicted speedup exceeds ×1.05. On held-out programs it achieves a mean "
        "speedup of **×1.21** against ×1.33 for the oracle, capturing **63%** of the achievable gain, while "
        "**9%** of its recommendations slow the program down; the speedup-prediction mean absolute error is "
        "0.105 (Figure 23, Table 8). Histogram Gradient Boosting was within 0.01 of it.", indent=True)
    R.figure(F("fig_reco.png"), "Held-out performance of the recommender compared with fixed policies", 4.6)
    R.table(["Strategy", "Model", "Reco. speedup", "Capture", "Regressions"], [
        ["argmax", "HistGradient Boosting", "×1.21", "63%", "14%"],
        ["argmax", "Random Forest", "×1.21", "64%", "12%"],
        ["per-pass classifiers", "HistGB classifier", "×1.22", "65%", "11%"],
        ["margin 0.03", "Random Forest", "×1.21", "64%", "10%"],
        ["**margin 0.05 (selected)**", "**Random Forest**", "**×1.21**", "**63%**", "**9%**"],
        ["margin 0.05", "HistGradient Boosting", "×1.21", "62%", "9%"],
        ["margin 0.08", "HistGradient Boosting", "×1.20", "60%", "6%"],
        ["margin 0.12", "Random Forest", "×1.20", "60%", "4%"],
        ["margin 0.18", "HistGradient Boosting", "×1.18", "53%", "2%"],
    ], "Selected rows of the cross-validated strategy bake-off (900 programs, 5 folds)",
        widths=[1.75, 1.8, 1.0, 0.8, 0.92], align=["l", "l", "c", "c", "c"])
    R.para(
        "Figure 24 shows the trade-off the margin controls. Raising it lowers the regression rate steeply "
        "while the captured gain falls only slowly, because most of the risky recommendations are those "
        "with small predicted gains. The features the model relies on most are the multiplication count, "
        "the estimated dynamic instruction count, the number of foldable constants, branch density, loop "
        "body size and instruction density inside loops — precisely the properties that determine whether "
        "CF, SR and LICM can pay off.", indent=True)
    R.figure(F("fig_margin.png"), "Effect of the abstain margin on regressions and captured gain", 6.0)
    R.h3("5.5.5 Model comparison")
    rows = [[m[0], m[1], f"×{m[2]:.3f}", f"{m[3]}%", f"{m[4]}%", f"{m[5]:.3f}", m[6]] for m in MODELS]
    rows[0] = [f"**{c}**" for c in rows[0]]
    R.table(["Model", "Input", "Reco.", "Capture", "Regress.", "MAE", "Size"], rows,
            "Eight-model comparison at each model's best margin under the 10% regression cap",
            widths=[1.65, 1.05, 0.7, 0.7, 0.7, 0.6, 0.87], align=["l", "l", "c", "c", "c", "c", "r"], size=9.5)
    R.para(
        "Histogram Gradient Boosting wins the comparison at ×1.209 and is also among the smallest models "
        "(2 MB, 0.36 ms to score all 64 combinations). The Random Forest scores almost the same but its "
        "serialised model is 613 MB. The most interesting result is the **1D-CNN**: reading only the raw "
        "TAC tokens, with none of the engineered features, it reaches ×1.197, within 1% of the best "
        "feature model (Figure 25).", indent=True)
    R.figure(F("fig_models.png"), "Recommendation quality of the eight models, with and without the "
                                  "regression cap", 6.0)
    R.para(
        "Without the cap, the **GRU** is actually the strongest model (×1.222, 67% capture), but 12% of its "
        "picks regress. Its predicted speedups are poorly calibrated: the gains it predicts for good and bad "
        "choices are similar, so no margin can filter the bad ones without discarding almost all the good "
        "ones, and under the cap it collapses to ×1.035. The hybrid models and the MLP show the same "
        "behaviour, which points to a training problem (calibration and data volume) rather than a lack of "
        "signal in their inputs.", indent=True)
    R.h3("5.5.6 Comparison with Review-1")
    R.table(["Measure", "Review-1 (270 programs)", "Final (900 programs)"], [
        ["Timed configurations", "17,280", "57,600"],
        ["Oracle geomean speedup", "×1.50", "×1.31"],
        ["All-on geomean speedup", "×1.21", "×1.19"],
        ["All-on not best", "243 / 270", "715 / 900"],
        ["Recommender held-out speedup", "×1.23 (HistGBT, margin 0.12)", "×1.21 (Random Forest, margin 0.05)"],
        ["Captured oracle gain", "44%", "63%"],
        ["Regressions (recommender / argmax)", "9% / 21%", "9% / 12%"],
        ["Models compared", "3 tree ensembles", "8 (incl. CNN, GRU, hybrids)"],
    ], "Results at Review-1 and in the final evaluation", widths=[2.3, 1.95, 2.02], align=["l", "c", "c"])
    R.para(
        "The oracle speedup fell because the 870 generated programs include branch-heavy, nested-loop and "
        "recursive patterns where the passes have less to work with; the larger corpus is more "
        "representative. On it, the recommender captures a much larger share of the available gain (63% "
        "against 44%) with the same 9% regression rate, which shows the value of the additional training "
        "data.", indent=True)
    R.h2("5.6 Testing and Verification")
    R.para(
        "Correctness is enforced at two levels. The pytest suite (49 tests plus 17 sub-tests) covers the "
        "lexer, parser, semantic analyser, IR generator, features, each pass individually and in all 64 "
        "combinations, regression cases for previously found miscompilations, nested aggregates, the timing "
        "harness and the ML data plumbing; the full suite passes. In addition, every one of the 57,600 "
        "builds in the experiment had to reproduce the unoptimised program's return value. Figure 26 shows "
        "the fast part of the suite.", indent=True)
    R.figure(F("term_pytest.png"), "Verbose run of the fast (non-compiling) part of the test suite: 25 tests pass", 5.6)

    # ===================================================================== 6
    R.h1("6. Conclusion", new_page=False)
    R.h2("6.1 Summary")
    R.para(
        "This project delivered a complete optimising compiler for MiniC — front end, three-address IR, "
        "control-flow analysis, six provably sound optimisation passes with a fixed-point pass manager, and "
        "a value-semantics C back end — together with a measurement harness, a 900-program benchmark corpus, "
        "a machine-learning pass recommender, a command-line tool and a browser playground.", indent=True)
    R.h2("6.2 Key Findings")
    R.bullets([
        "Pass selection matters and is program-specific: the best combination gives ×1.31 on average and up "
        "to ×4.37, but enabling everything is sub-optimal for 79% of programs and harmful for 12%.",
        "Every fixed non-empty combination slows at least one program down, so a per-program decision is "
        "required to be both fast and safe.",
        "A Random Forest over 27 static features, with an abstain margin, achieves ×1.21 on unseen programs "
        "and captures 63% of the achievable gain with only 9% regressions.",
        "Risk control is as important as accuracy: the abstain margin cut regressions from 12% to 9% at "
        "almost no cost in speedup.",
        "Tree ensembles over engineered features perform best, but a CNN over the raw TAC sequence comes "
        "within 1% of them, showing that useful program representations can be learned directly from IR.",
    ])
    R.h2("6.3 Limitations")
    R.bullets([
        "Speedups are measured with `gcc -O0` to isolate the project's passes; the gains on top of an "
        "optimising C compiler would be smaller.",
        "All timings come from one machine; absolute speedups depend on the processor and memory system.",
        "The passes are intra-procedural and their order is fixed, so recursive programs and phase ordering "
        "are only partly addressed.",
        "The generated programs, although varied, follow known patterns and may be easier to predict than "
        "arbitrary real-world code.",
        "The selected Random Forest is large (613 MB serialised); the comparison shows a 2 MB gradient "
        "boosting model performs at least as well.",
    ])
    R.h2("6.4 Future Work")
    R.bullets([
        "Ship Histogram Gradient Boosting as the default recommender for its size and speed.",
        "Calibrate the sequence models (for example with quantile or ranking losses) so that the GRU's "
        "uncapped advantage can be used safely.",
        "Treat pass ordering as a sequential decision problem and apply reinforcement learning, as in "
        "CompilerGym and AutoPhase [4], [5].",
        "Add inter-procedural optimisation, in particular function inlining, to help recursive and "
        "call-heavy programs.",
        "Augment static features with hardware performance counters, and target LLVM IR to evaluate the "
        "approach on top of a production optimiser.",
    ])

    # ===================================================================== 7
    R.h1("7. Project Repository Link", new_page=False)
    R.para("The complete source code, benchmark corpus, measured dataset, trained-model pipeline, "
           "results and this report's generator are available at:", indent=False)
    p = R.para("**https://github.com/AkineHime/AI-based-compiler-optimization**",
               align=WD_ALIGN_PARAGRAPH.CENTER)
    p.paragraph_format.space_after = Pt(12)
    R.para("Key files: `RESULTS.md` (measured results), `MODEL_COMPARISON.md` (eight-model comparison), "
           "`docs/WALKTHROUGH.md` (one program traced through every stage), `docs/results.html` (visual "
           "report) and `run_experiment.py` (reproduces the whole experiment).")

    # ===================================================================== refs
    R.h1("References", new_page=False)
    refs = [
        "A. V. Aho, M. S. Lam, R. Sethi and J. D. Ullman, _Compilers: Principles, Techniques, and Tools_, "
        "2nd ed. Boston, MA, USA: Addison-Wesley, 2006.",
        "K. D. Cooper and L. Torczon, _Engineering a Compiler_, 2nd ed. San Francisco, CA, USA: Morgan "
        "Kaufmann, 2011.",
        "G. Fursin et al., \"Milepost GCC: Machine learning enabled self-tuning compiler,\" _International "
        "Journal of Parallel Programming_, vol. 39, no. 3, pp. 296–327, 2011.",
        "C. Cummins et al., \"CompilerGym: Robust, performant compiler optimization environments for AI "
        "research,\" in _Proc. IEEE/ACM Int. Symp. Code Generation and Optimization (CGO)_, 2022.",
        "A. Haj-Ali et al., \"AutoPhase: Juggling HLS phase orderings in random forests with deep "
        "reinforcement learning,\" in _Proc. Machine Learning and Systems (MLSys)_, 2020.",
        "M. Trofin, Y. Qian, E. Brevdo, Z. Lin, K. Choromanski and D. Li, \"MLGO: A machine learning guided "
        "compiler optimizations framework,\" arXiv:2101.04808, 2021.",
        "A. H. Ashouri, W. Killian, J. Cavazos, G. Palermo and C. Silvano, \"A survey on compiler autotuning "
        "using machine learning,\" _ACM Computing Surveys_, vol. 51, no. 5, 2018.",
        "F. Pedregosa et al., \"Scikit-learn: Machine learning in Python,\" _Journal of Machine Learning "
        "Research_, vol. 12, pp. 2825–2830, 2011.",
        "A. Paszke et al., \"PyTorch: An imperative style, high-performance deep learning library,\" in "
        "_Advances in Neural Information Processing Systems 32 (NeurIPS)_, 2019.",
        "G. Ke et al., \"LightGBM: A highly efficient gradient boosting decision tree,\" in _Advances in "
        "Neural Information Processing Systems 30 (NIPS)_, 2017.",
        "L. Breiman, \"Random forests,\" _Machine Learning_, vol. 45, no. 1, pp. 5–32, 2001.",
        "S. S. Muchnick, _Advanced Compiler Design and Implementation_. San Francisco, CA, USA: Morgan "
        "Kaufmann, 1997.",
    ]
    for i, r in enumerate(refs, 1):
        p = R.d.add_paragraph()
        pf = p.paragraph_format
        pf.left_indent = Cm(1.0)
        pf.first_line_indent = Cm(-1.0)
        pf.space_after = Pt(6)
        pf.line_spacing = 1.15
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        _font(p.add_run(f"[{i}]\t"), size=11)
        pf.tab_stops.add_tab_stop(Cm(1.0))
        R._runs(p, r, size=11)


MEMBERS = [("24BCT0301", "Ashmit Verma"), ("24BAI0243", "S Karthik Sudharsan"),
           ("23BCE2091", "Mohammed Nadeem")]


def build(pages=None, out=OUT):
    R = Report(TEMPLATE)
    R.title_page("AI-Based Compiler Optimization for MiniC",
                 "Static-Feature-Driven Pass Selection Using Machine Learning", MEMBERS, "November 2026")
    R.page_numbers()
    R.page_break()
    R.contents(pages)
    write_body(R)
    R.save(out)
    return R


def find_pages(pdf_path, R):
    """Page number (1-based, as printed in the footer) of every heading, figure and table."""
    import pymupdf
    doc = pymupdf.open(pdf_path)
    texts = [" ".join(pg.get_text().split()) for pg in doc]
    start = next(i for i, t in enumerate(texts) if "1. Abstract" in t and "TABLE OF CONTENTS" not in t
                 and "LIST OF" not in t)
    pages = {}

    def first(needle):
        for i in range(start, len(texts)):
            if needle in texts[i]:
                return i + 1
        return None

    for level, text in R.heads:
        pages[("h", text)] = first(" ".join(text.split()))
    for n, cap in R.figs:
        pages[("f", n)] = first(f"Figure {n}:")
    for n, cap in R.tabs:
        pages[("t", n)] = first(f"Table {n}:")
    return pages, len(doc)


def to_pdf(docx_path):
    import subprocess
    soffice = r"C:\Program Files\LibreOffice\program\soffice.exe"
    outdir = os.path.dirname(docx_path)
    subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", outdir, docx_path],
                   check=True, capture_output=True)
    return os.path.splitext(docx_path)[0] + ".pdf"


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--no-figures", action="store_true")
    ap.add_argument("--no-pdf", action="store_true", help="skip the LibreOffice page-number pass")
    args = ap.parse_args()
    if not args.no_figures:
        figures()
        print("figures ->", FIG)
    # pass 1: collect headings / figures / tables
    R = build()
    CONTENTS_HEADINGS[:] = R.heads
    FIGURE_LIST[:] = R.figs
    TABLE_LIST[:] = R.tabs
    # pass 2: same layout with placeholder page numbers -> render -> real numbers
    build()
    if args.no_pdf:
        print("wrote", OUT, "(contents page numbers are placeholders)")
        return
    pages, n = find_pages(to_pdf(OUT), R)
    missing = [k for k, v in pages.items() if v is None]
    if missing:
        print("WARNING: not found in PDF:", missing[:5])
    build(pages)
    pdf = to_pdf(OUT)
    import pymupdf
    print(f"wrote {OUT}\nwrote {pdf}  ({len(pymupdf.open(pdf))} pages)")


if __name__ == "__main__":
    main()
